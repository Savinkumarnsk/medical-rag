"""
loaders.py — Document loaders for PDF, DOCX, TXT, MD.
Extracts text with metadata: source, page, section.
"""
from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import List

from langchain_core.documents import Document

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────────────────────────────────────
# PDF Loader (pymupdf4llm with OCR fallback)
# ─────────────────────────────────────────────────────────────────────────────

def load_pdf(file_path: Path) -> List[Document]:
    """Load PDF with pymupdf4llm; fall back to pytesseract OCR for scanned pages."""
    docs: List[Document] = []

    try:
        import pymupdf4llm
        md_text = pymupdf4llm.to_markdown(str(file_path), page_chunks=True)

        for page_data in md_text:
            text = page_data.get("text", "") if isinstance(page_data, dict) else str(page_data)
            page_num = page_data.get("metadata", {}).get("page", 0) + 1 if isinstance(page_data, dict) else 0
            section = _extract_section_heading(text)

            # If page is essentially empty, attempt OCR
            if len(text.strip()) < 50:
                ocr_text = _ocr_page(file_path, page_num - 1)
                if ocr_text:
                    text = ocr_text

            if text.strip():
                docs.append(Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "source_path": str(file_path),
                        "page": page_num,
                        "section": section,
                        "file_type": "pdf",
                    }
                ))
    except Exception as exc:
        logger.warning("pymupdf4llm failed for %s: %s. Falling back to PyMuPDF.", file_path.name, exc)
        docs = _load_pdf_fallback(file_path)

    logger.info("Loaded %d pages from %s", len(docs), file_path.name)
    return docs


def _load_pdf_fallback(file_path: Path) -> List[Document]:
    """Basic PyMuPDF page-by-page fallback."""
    import fitz  # PyMuPDF
    docs = []
    with fitz.open(str(file_path)) as pdf:
        for page_num, page in enumerate(pdf, start=1):
            text = page.get_text("text")
            if len(text.strip()) < 50:
                text = _ocr_page(file_path, page_num - 1) or text
            if text.strip():
                docs.append(Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "source_path": str(file_path),
                        "page": page_num,
                        "section": _extract_section_heading(text),
                        "file_type": "pdf",
                    }
                ))
    return docs


def _ocr_page(file_path: Path, page_index: int) -> str:
    """Run pytesseract OCR on a specific PDF page."""
    try:
        import fitz
        from PIL import Image
        import pytesseract
        import io

        with fitz.open(str(file_path)) as pdf:
            page = pdf[page_index]
            pix = page.get_pixmap(dpi=300)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            return pytesseract.image_to_string(img)
    except Exception as exc:
        logger.debug("OCR failed on page %d of %s: %s", page_index, file_path.name, exc)
        return ""


# ─────────────────────────────────────────────────────────────────────────────
# DOCX Loader
# ─────────────────────────────────────────────────────────────────────────────

def load_docx(file_path: Path) -> List[Document]:
    """Load DOCX, extracting paragraphs grouped by sections."""
    from docx import Document as DocxDocument  # type: ignore

    doc = DocxDocument(str(file_path))
    docs: List[Document] = []
    current_section = "Introduction"
    buffer: List[str] = []
    pseudo_page = 1

    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            if buffer:
                docs.append(Document(
                    page_content="\n".join(buffer),
                    metadata={
                        "source": file_path.name,
                        "source_path": str(file_path),
                        "page": pseudo_page,
                        "section": current_section,
                        "file_type": "docx",
                    }
                ))
                buffer = []
                pseudo_page += 1
            current_section = para.text.strip()
        else:
            text = para.text.strip()
            if text:
                buffer.append(text)

    if buffer:
        docs.append(Document(
            page_content="\n".join(buffer),
            metadata={
                "source": file_path.name,
                "source_path": str(file_path),
                "page": pseudo_page,
                "section": current_section,
                "file_type": "docx",
            }
        ))

    logger.info("Loaded %d sections from %s", len(docs), file_path.name)
    return docs


# ─────────────────────────────────────────────────────────────────────────────
# TXT Loader
# ─────────────────────────────────────────────────────────────────────────────

def load_txt(file_path: Path) -> List[Document]:
    """Load plain text file as a single document."""
    text = file_path.read_text(encoding="utf-8", errors="replace")
    return [Document(
        page_content=text,
        metadata={
            "source": file_path.name,
            "source_path": str(file_path),
            "page": 1,
            "section": "Document",
            "file_type": "txt",
        }
    )]


# ─────────────────────────────────────────────────────────────────────────────
# Markdown Loader
# ─────────────────────────────────────────────────────────────────────────────

def load_md(file_path: Path) -> List[Document]:
    """Load Markdown, splitting by headers to produce section-aware docs."""
    from langchain_community.document_loaders import UnstructuredMarkdownLoader

    text = file_path.read_text(encoding="utf-8", errors="replace")
    sections = re.split(r"(?m)^#{1,3} .+$", text)
    headings = re.findall(r"(?m)^#{1,3} .+$", text)

    docs: List[Document] = []
    for i, section_text in enumerate(sections):
        if not section_text.strip():
            continue
        heading = headings[i - 1].lstrip("#").strip() if i > 0 and i - 1 < len(headings) else "Preamble"
        docs.append(Document(
            page_content=section_text.strip(),
            metadata={
                "source": file_path.name,
                "source_path": str(file_path),
                "page": i + 1,
                "section": heading,
                "file_type": "md",
            }
        ))

    if not docs:
        docs = [Document(
            page_content=text,
            metadata={
                "source": file_path.name,
                "source_path": str(file_path),
                "page": 1,
                "section": "Document",
                "file_type": "md",
            }
        )]

    logger.info("Loaded %d sections from %s", len(docs), file_path.name)
    return docs


# ─────────────────────────────────────────────────────────────────────────────
# Dispatcher
# ─────────────────────────────────────────────────────────────────────────────

LOADER_MAP = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
    ".md": load_md,
}


def load_document(file_path: Path) -> List[Document]:
    """Load a document based on its extension."""
    ext = file_path.suffix.lower()
    loader_fn = LOADER_MAP.get(ext)
    if loader_fn is None:
        raise ValueError(f"Unsupported file type: {ext}")
    return loader_fn(file_path)


def load_directory(dir_path: Path) -> List[Document]:
    """Load all supported documents from a directory recursively."""
    all_docs: List[Document] = []
    for ext in LOADER_MAP:
        for fp in sorted(dir_path.rglob(f"*{ext}")):
            try:
                docs = load_document(fp)
                all_docs.extend(docs)
            except Exception as exc:
                logger.error("Failed to load %s: %s", fp, exc)
    return all_docs


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _extract_section_heading(text: str) -> str:
    """Extract the first markdown or all-caps heading from text."""
    # Markdown heading
    m = re.search(r"^#{1,4}\s+(.+)$", text, re.MULTILINE)
    if m:
        return m.group(1).strip()
    # ALL CAPS line (common in medical PDFs)
    m = re.search(r"^([A-Z][A-Z\s]{4,60})$", text, re.MULTILINE)
    if m:
        return m.group(1).strip().title()
    return "General"
